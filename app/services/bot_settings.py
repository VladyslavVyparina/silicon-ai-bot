"""Налаштування бота (модель GPT, промпт, база знань) — лише для owner у адмінці."""
from __future__ import annotations

import json
import re
from pathlib import Path

from sqlalchemy import select
from sqlalchemy.ext.asyncio import AsyncSession

from app.config import PROJECT_ROOT, get_settings
from app.database.models import AppSetting

KNOWLEDGE_DIR = PROJECT_ROOT / "data" / "knowledge"
PROMPT_OVERRIDE_PATH = PROJECT_ROOT / "data" / "consultant_prompt_override.txt"
SETTINGS_CACHE_PATH = PROJECT_ROOT / "data" / "bot_runtime.json"

OPENAI_CHAT_MODELS = [
    "gpt-5.6",
    "gpt-5.6-sol",
    "gpt-5.6-terra",
    "gpt-5.6-luna",
    "gpt-5.5",
    "gpt-5.5-pro",
    "gpt-5.4",
    "gpt-5.4-mini",
    "gpt-5.4-nano",
    "gpt-5.2",
    "gpt-5.1",
    "gpt-5",
    "gpt-5-mini",
    "gpt-5-nano",
    "gpt-5-pro",
    "gpt-4.1",
    "gpt-4.1-mini",
    "gpt-4.1-nano",
    "gpt-4o",
    "gpt-4o-mini",
    "o3",
    "o3-mini",
    "o4-mini",
]

KEY_MODEL = "openai_model"
KEY_PROMPT = "consultant_prompt"


def ensure_dirs() -> None:
    KNOWLEDGE_DIR.mkdir(parents=True, exist_ok=True)
    (PROJECT_ROOT / "data").mkdir(parents=True, exist_ok=True)


async def get_setting(db: AsyncSession, key: str, default: str = "") -> str:
    row = (
        await db.execute(select(AppSetting).where(AppSetting.key == key))
    ).scalar_one_or_none()
    if row and row.value is not None:
        return row.value
    return default


async def set_setting(db: AsyncSession, key: str, value: str) -> None:
    row = (
        await db.execute(select(AppSetting).where(AppSetting.key == key))
    ).scalar_one_or_none()
    if row:
        row.value = value
    else:
        db.add(AppSetting(key=key, value=value))
    await db.commit()
    _write_runtime_cache(key, value)


def _write_runtime_cache(key: str, value: str) -> None:
    ensure_dirs()
    data: dict = {}
    if SETTINGS_CACHE_PATH.exists():
        try:
            data = json.loads(SETTINGS_CACHE_PATH.read_text(encoding="utf-8"))
        except Exception:
            data = {}
    data[key] = value
    SETTINGS_CACHE_PATH.write_text(json.dumps(data, ensure_ascii=False, indent=2), encoding="utf-8")


def _read_runtime_cache(key: str) -> str | None:
    if not SETTINGS_CACHE_PATH.exists():
        return None
    try:
        data = json.loads(SETTINGS_CACHE_PATH.read_text(encoding="utf-8"))
        val = data.get(key)
        return str(val) if val is not None else None
    except Exception:
        return None


def get_openai_model_sync() -> str:
    cached = _read_runtime_cache(KEY_MODEL)
    if cached:
        return cached
    return get_settings().openai_model or "gpt-5"


DEFAULT_GREETING = (
    "Вітаю! Мене звати Олег, я менеджер компанії Silicon. "
    "Чим можу вам допомогти?"
)


def get_prompt_sync() -> str:
    """Промпт: свіжий файл інструкцій / override з правилом «питання першим».

    Довгий застарілий override з адмінки (без НА ПЕРШОМУ МІСЦІ) ігноруємо —
    інакше GPT знову крутить воронку і «не чує» питання. Короткі кастомні
    промпти з адмінки (тести / експерименти) залишаємо як є.
    """
    repo = ""
    path = Path(get_settings().consultant_instructions_path)
    if not path.is_absolute():
        path = PROJECT_ROOT / path
    if path.exists():
        repo = path.read_text(encoding="utf-8")

    def _is_fresh(text: str) -> bool:
        t = (text or "").strip()
        if len(t) < 40:
            return False
        return (
            "НА ПЕРШОМУ МІСЦІ" in t
            or "Спочатку відповіси на останнє повідомлення" in t
            or "Спочатку зрозумій і відповіси" in t
        )

    def _is_stale_long(text: str) -> bool:
        t = (text or "").strip()
        # довгий «старий» промпт без правила питання-першим
        return len(t) >= 2000 and not _is_fresh(t)

    if PROMPT_OVERRIDE_PATH.exists():
        text = PROMPT_OVERRIDE_PATH.read_text(encoding="utf-8").strip()
        if text and not _is_stale_long(text):
            return text
        if text and _is_stale_long(text) and repo and _is_fresh(repo):
            return repo
        if text:
            return text
    cached = _read_runtime_cache(KEY_PROMPT)
    if cached and not _is_stale_long(cached):
        return cached.strip()
    if repo:
        return repo
    if cached and len(cached.strip()) > 40:
        return cached.strip()
    return ""


def get_greeting_sync() -> str:
    """Вітання з промпту адмінки (блок ВІТАННЯ), інакше дефолт."""
    prompt = get_prompt_sync()
    m = re.search(r"ВІТАННЯ.{0,500}[«\"]([^»\"]{15,240})[»\"]", prompt, re.S)
    if m:
        return m.group(1).strip()
    return DEFAULT_GREETING


async def restore_prompt_override_from_db(db: AsyncSession) -> None:
    """Якщо деплой стер файл адмінки — підняти промпт з БД."""
    ensure_dirs()
    if PROMPT_OVERRIDE_PATH.exists() and PROMPT_OVERRIDE_PATH.read_text(encoding="utf-8").strip():
        return
    stored = await get_setting(db, KEY_PROMPT, "")
    if stored and len(stored.strip()) > 40:
        PROMPT_OVERRIDE_PATH.write_text(stored, encoding="utf-8")
        _write_runtime_cache(KEY_PROMPT, stored)


async def save_prompt(db: AsyncSession, text: str) -> None:
    ensure_dirs()
    PROMPT_OVERRIDE_PATH.write_text(text or "", encoding="utf-8")
    await set_setting(db, KEY_PROMPT, text or "")


async def save_model(db: AsyncSession, model: str) -> None:
    model = (model or "").strip()
    if model not in OPENAI_CHAT_MODELS:
        raise ValueError(f"Невідома модель: {model}")
    await set_setting(db, KEY_MODEL, model)


def list_knowledge_files() -> list[dict]:
    ensure_dirs()
    files = []
    for p in sorted(KNOWLEDGE_DIR.iterdir(), key=lambda x: x.name.lower()):
        if p.is_file() and not p.name.startswith("."):
            try:
                rel = str(p.relative_to(PROJECT_ROOT))
            except ValueError:
                rel = str(p)
            files.append({
                "name": p.name,
                "size": p.stat().st_size,
                "path": rel,
            })
    return files


def save_knowledge_file(filename: str, content: bytes) -> str:
    ensure_dirs()
    safe = "".join(c for c in filename if c.isalnum() or c in ("-", "_", ".", " ")).strip()
    if not safe:
        safe = "knowledge.txt"
    dest = KNOWLEDGE_DIR / safe
    dest.write_bytes(content)
    return safe


def delete_knowledge_file(filename: str) -> bool:
    ensure_dirs()
    safe = Path(filename).name
    dest = KNOWLEDGE_DIR / safe
    if dest.exists() and dest.is_file():
        dest.unlink()
        return True
    return False


def _looks_like_pdf_garbage(text: str) -> bool:
    sample = text[:4000]
    return sample.count("/uni") >= 8 or sample.count("uni0") >= 12


def _read_pdf_text(path: Path) -> str:
    text = ""
    try:
        import fitz  # pymupdf

        doc = fitz.open(str(path))
        parts = [(page.get_text() or "") for page in doc[:40]]
        doc.close()
        text = "\n".join(parts)
    except Exception:
        text = ""
    if not text.strip() or _looks_like_pdf_garbage(text):
        try:
            from pypdf import PdfReader

            reader = PdfReader(str(path))
            parts = [(page.extract_text() or "") for page in reader.pages[:40]]
            text = "\n".join(parts)
        except Exception:
            return ""
    if _looks_like_pdf_garbage(text):
        return ""
    text = re.sub(r"[ \t]+\n", "\n", text)
    text = re.sub(r"\n{3,}", "\n\n", text)
    return text.strip()


def _read_knowledge_file(path: Path) -> str:
    suffix = path.suffix.lower()
    if suffix in {".txt", ".md", ".csv", ".json"}:
        try:
            return path.read_text(encoding="utf-8", errors="ignore").strip()
        except Exception:
            return ""
    if suffix == ".docx":
        try:
            from docx import Document

            doc = Document(str(path))
            parts = [p.text.strip() for p in doc.paragraphs if p.text and p.text.strip()]
            for table in doc.tables:
                for row in table.rows:
                    cells = [c.text.strip() for c in row.cells if c.text.strip()]
                    if cells:
                        parts.append(" | ".join(cells))
            return "\n".join(parts).strip()
        except Exception:
            return ""
    if suffix == ".pdf":
        return _read_pdf_text(path)
    return ""


def load_knowledge_text(limit_chars: int = 80000) -> str:
    """Прайс з data/knowledge + повна база app/data/consultant_knowledge_full.txt."""
    ensure_dirs()
    files = [p for p in KNOWLEDGE_DIR.iterdir() if p.is_file() and not p.name.startswith(".")]
    has_readable = any(p.suffix.lower() in {".txt", ".md", ".docx"} for p in files)
    chunks: list[str] = []
    total = 0
    ext_rank = {".txt": 0, ".md": 1, ".docx": 2, ".csv": 3, ".json": 4, ".pdf": 5}

    def sort_key(p: Path):
        return (ext_rank.get(p.suffix.lower(), 9), p.name.lower())

    def _append(text: str, label: str = "") -> None:
        nonlocal total
        if not text or total >= limit_chars:
            return
        block = f"--- {label} ---\n{text}\n" if label else f"---\n{text}\n"
        if total + len(block) > limit_chars:
            remain = limit_chars - total
            if remain > 100:
                chunks.append(block[:remain])
                total = limit_chars
            return
        chunks.append(block)
        total += len(block)

    for p in sorted(files, key=sort_key):
        suffix = p.suffix.lower()
        if suffix not in {".txt", ".md", ".csv", ".json", ".pdf", ".docx"}:
            continue
        if suffix == ".pdf" and has_readable:
            continue
        _append(_read_knowledge_file(p), p.name)

    full_path = PROJECT_ROOT / "app" / "data" / "consultant_knowledge_full.txt"
    if full_path.exists():
        try:
            full_text = full_path.read_text(encoding="utf-8", errors="ignore").strip()
        except Exception:
            full_text = ""
        if full_text:
            _append(full_text, "consultant_knowledge_full.txt")

    return "\n".join(chunks).strip()
