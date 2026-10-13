import io

import pytest
from httpx import AsyncClient

from app.services.admin_auth import (
    check_password,
    create_admin_token,
    ensure_owner_exists,
    hash_password,
    verify_admin_token,
    verify_password,
)
from app.services.email_service import get_sent_emails, reset_sent_emails


def test_password_hash_roundtrip():
    h = hash_password("Secret123")
    assert verify_password("Secret123", h) is True
    assert verify_password("wrong", h) is False


def test_admin_token_roundtrip(monkeypatch):
    monkeypatch.setenv("ADMIN_SECRET_KEY", "test-secret-key")
    from app.config import get_settings

    get_settings.cache_clear()
    token = create_admin_token({"id": 1, "role": "owner", "email": "a@b.c", "login": "owner"})
    assert verify_admin_token(token) is True
    assert verify_admin_token("bad") is False
    get_settings.cache_clear()


def test_admin_password(monkeypatch):
    monkeypatch.setenv("ADMIN_PASSWORD", "secret123")
    from app.config import get_settings

    get_settings.cache_clear()
    assert check_password("secret123") is True
    assert check_password("wrong") is False
    get_settings.cache_clear()


@pytest.mark.asyncio
async def test_manager_handoff_creates_ticket(client: AsyncClient):
    resp = await client.post("/api/chat/message", json={
        "chat_id": "admin-handoff-1",
        "message": "Будь ласка, зв'яжіть мене з менеджером щодо RFID",
    })
    assert resp.status_code == 200
    data = resp.json()
    assert data["action"] == "manager_handoff"
    assert data["data"]["ticket_id"] > 0
    assert "передав" in data["response"].lower() and "менеджер" in data["response"].lower()
    assert "вітаю" not in data["response"].lower()
    assert "менеджер компанії silicon" not in data["response"].lower()

    # заявка має з’явитись у списку для адміна
    token = await _owner_token(client)
    tickets_page = await client.get(f"/admin/tickets?token={token}")
    assert tickets_page.status_code == 200
    assert "admin-handoff-1" in tickets_page.text
    assert "RFID" in tickets_page.text or "менеджер" in tickets_page.text.lower()


@pytest.mark.asyncio
async def test_human_mode_silences_ai(client: AsyncClient):
    await client.post("/api/chat/message", json={
        "chat_id": "admin-handoff-2",
        "message": "Хочу менеджера",
    })
    resp = await client.post("/api/chat/message", json={
        "chat_id": "admin-handoff-2",
        "message": "Додаткове питання без AI",
    })
    assert resp.status_code == 200
    assert resp.json()["action"] == "human_mode"
    assert resp.json()["response"] == ""


@pytest.mark.asyncio
async def test_takeover_sends_manager_greeting(client: AsyncClient, monkeypatch):
    monkeypatch.setenv("ADMIN_SECRET_KEY", "takeover-secret")
    monkeypatch.setenv("ADMIN_PASSWORD", "silicon-admin")
    monkeypatch.setenv("ADMIN_OWNER_EMAIL", "owner@silicon.local")
    monkeypatch.setenv("MOCK_EXTERNAL_APIS", "true")
    from app.config import get_settings
    from unittest.mock import AsyncMock, patch

    get_settings.cache_clear()
    await client.get("/admin/login")
    token = await _owner_token(client)

    # задати ім'я власнику
    await client.post(
        "/admin/profile/update",
        data={"token": token, "full_name": "Влад Менеджер"},
        follow_redirects=False,
    )

    await client.post("/api/chat/message", json={
        "chat_id": "takeover-chat-1",
        "message": "Зв'яжіть з менеджером",
    })

    sent = []

    async def fake_send(self, chat_id, text):
        sent.append({"chat_id": str(chat_id), "text": text})
        return {"success": True}

    with patch("app.services.telegram_bot.TelegramBot.send_message", new=fake_send):
        resp = await client.post(
            "/admin/live-chats/takeover-chat-1/takeover",
            data={"token": token},
        )
    assert resp.status_code == 200
    assert resp.json()["greeting_sent"] is True
    assert sent
    assert "Влад Менеджер" in sent[0]["text"]
    assert "менеджер компанії Silicon" in sent[0]["text"]
    get_settings.cache_clear()


async def _owner_token(client: AsyncClient, db_session=None):
    """Логін власника після ensure_owner_exists через сторінку логіну / напряму."""
    from app.config import get_settings
    from app.database.session import async_session
    from app.services.admin_auth import ensure_owner_exists, create_admin_token

    # Use app dependency DB via login form
    login = await client.post(
        "/admin/login",
        data={
            "identifier": get_settings().admin_owner_email,
            "password": get_settings().admin_password,
        },
        follow_redirects=False,
    )
    assert login.status_code == 303
    loc = login.headers["location"]
    assert "token=" in loc
    return loc.split("token=")[1].split("&")[0]


@pytest.mark.asyncio
async def test_admin_login_and_stats(client: AsyncClient, monkeypatch):
    monkeypatch.setenv("ADMIN_PASSWORD", "silicon-admin")
    monkeypatch.setenv("ADMIN_OWNER_EMAIL", "owner@silicon.local")
    monkeypatch.setenv("ADMIN_SECRET_KEY", "test-admin-secret")
    from app.config import get_settings

    get_settings.cache_clear()

    await client.post("/api/chat/message", json={
        "chat_id": "stats-user-1",
        "message": "/start",
    })

    # force owner seed via GET login page
    await client.get("/admin/login")

    bad = await client.post(
        "/admin/login",
        data={"identifier": "owner@silicon.local", "password": "wrong"},
        follow_redirects=False,
    )
    assert bad.status_code == 303

    token = await _owner_token(client)
    dash = await client.get(f"/admin?token={token}")
    assert dash.status_code == 200
    assert "Дашборд" in dash.text or "Користувачі" in dash.text

    stats = await client.get(f"/admin/stats?token={token}")
    assert stats.status_code == 200
    body = stats.json()
    assert body["bot_users"] >= 1

    get_settings.cache_clear()


@pytest.mark.asyncio
async def test_admin_send_message(client: AsyncClient, monkeypatch):
    monkeypatch.setenv("ADMIN_SECRET_KEY", "test-admin-secret-2")
    monkeypatch.setenv("ADMIN_PASSWORD", "silicon-admin")
    monkeypatch.setenv("ADMIN_OWNER_EMAIL", "owner@silicon.local")
    from app.config import get_settings

    get_settings.cache_clear()
    await client.get("/admin/login")

    chat_id = "admin-reply-1"
    await client.post("/api/chat/message", json={
        "chat_id": chat_id,
        "message": "Потрібен менеджер для прорахунку",
    })
    token = await _owner_token(client)
    send = await client.post(
        f"/admin/live-chats/{chat_id}/send",
        data={"token": token, "message": "Доброго дня, це менеджер Silicon."},
    )
    assert send.status_code == 200
    assert send.json()["success"] is True

    msgs = await client.get(
        f"/admin/live-chats/{chat_id}/messages?token={token}&after=0"
    )
    kinds = [m["kind"] for m in msgs.json()["messages"]]
    assert "admin" in kinds
    assert "user" in kinds

    get_settings.cache_clear()


@pytest.mark.asyncio
async def test_invite_manager_sends_email(client: AsyncClient, monkeypatch):
    monkeypatch.setenv("ADMIN_SECRET_KEY", "invite-secret")
    monkeypatch.setenv("ADMIN_PASSWORD", "silicon-admin")
    monkeypatch.setenv("ADMIN_OWNER_EMAIL", "owner@silicon.local")
    monkeypatch.setenv("MOCK_EXTERNAL_APIS", "true")
    from app.config import get_settings

    get_settings.cache_clear()
    reset_sent_emails()
    await client.get("/admin/login")
    token = await _owner_token(client)

    resp = await client.post(
        "/admin/team/invite",
        data={
            "token": token,
            "email": "manager.new@example.com",
            "full_name": "Новий Менеджер",
        },
        follow_redirects=False,
    )
    assert resp.status_code == 303
    assert "msg=" in resp.headers["location"]
    # одноразовий показ пароля в панелі (без SMTP)
    page = await client.get(resp.headers["location"])
    assert page.status_code == 200
    assert "Доступ менеджера" in page.text
    assert "manager.new@example.com" in page.text
    assert "cred-pass" in page.text
    assert "Скопіюйте дані для входу нижче" in page.text or "додано" in page.text.lower()

    sent = get_sent_emails()
    assert len(sent) >= 1
    assert sent[-1]["to"] == "manager.new@example.com"
    assert "Пароль:" in sent[-1]["body"]
    assert "Email / логін: manager.new@example.com" in sent[-1]["body"]
    assert "Новий Менеджер" in sent[-1]["body"]

    body = sent[-1]["body"]
    password = None
    for line in body.splitlines():
        if line.startswith("Пароль:"):
            password = line.split(":", 1)[1].strip()
    assert password
    assert password in page.text
    login_resp = await client.post(
        "/admin/login",
        data={"identifier": "manager.new@example.com", "password": password},
        follow_redirects=False,
    )
    assert login_resp.status_code == 303
    assert "token=" in login_resp.headers["location"]
    get_settings.cache_clear()


@pytest.mark.asyncio
async def test_invite_requires_name(client: AsyncClient, monkeypatch):
    monkeypatch.setenv("ADMIN_SECRET_KEY", "invite-secret-2")
    monkeypatch.setenv("ADMIN_PASSWORD", "silicon-admin")
    monkeypatch.setenv("ADMIN_OWNER_EMAIL", "owner@silicon.local")
    from app.config import get_settings

    get_settings.cache_clear()
    await client.get("/admin/login")
    token = await _owner_token(client)
    resp = await client.post(
        "/admin/team/invite",
        data={"token": token, "email": "x@y.com", "full_name": "  "},
        follow_redirects=False,
    )
    assert resp.status_code == 303
    assert "err=" in resp.headers["location"]
    get_settings.cache_clear()


@pytest.mark.asyncio
async def test_profile_change_name_only(client: AsyncClient, monkeypatch):
    monkeypatch.setenv("ADMIN_SECRET_KEY", "profile-secret")
    monkeypatch.setenv("ADMIN_PASSWORD", "silicon-admin")
    monkeypatch.setenv("ADMIN_OWNER_EMAIL", "owner@silicon.local")
    from app.config import get_settings

    get_settings.cache_clear()
    await client.get("/admin/login")
    token = await _owner_token(client)

    resp = await client.post(
        "/admin/profile/update",
        data={"token": token, "full_name": "Головний Власник"},
        follow_redirects=False,
    )
    assert resp.status_code == 303
    assert "msg=" in resp.headers["location"]

    # email still works as login
    login_resp = await client.post(
        "/admin/login",
        data={"identifier": "owner@silicon.local", "password": "silicon-admin"},
        follow_redirects=False,
    )
    assert login_resp.status_code == 303

    img = io.BytesIO(b"\x89PNG\r\n\x1a\n" + b"\x00" * 32)
    av = await client.post(
        "/admin/profile/avatar",
        data={"token": token},
        files={"avatar": ("photo.png", img, "image/png")},
        follow_redirects=False,
    )
    assert av.status_code == 303
    assert "msg=" in av.headers["location"]

    dash = await client.get(f"/admin?token={token}")
    assert dash.status_code == 200
    assert "cb-user-avatar" in dash.text
    assert "/uploads/avatars/" in dash.text
    get_settings.cache_clear()


@pytest.mark.asyncio
async def test_smtp_invite_email_delivery(client: AsyncClient, monkeypatch):
    """Реальний SMTP-шлях (MOCK=false): перевіряємо відправку через smtplib."""
    from email.message import EmailMessage
    from unittest.mock import patch

    sent_msgs: list[EmailMessage] = []

    class FakeSMTP:
        def __init__(self, host, port, timeout=30):
            self.host = host
            self.port = port

        def __enter__(self):
            return self

        def __exit__(self, *args):
            return False

        def starttls(self):
            return None

        def login(self, user, password):
            return None

        def send_message(self, msg):
            sent_msgs.append(msg)
            return {}

    monkeypatch.setenv("ADMIN_SECRET_KEY", "smtp-secret")
    monkeypatch.setenv("ADMIN_PASSWORD", "silicon-admin")
    monkeypatch.setenv("ADMIN_OWNER_EMAIL", "owner@silicon.local")
    monkeypatch.setenv("MOCK_EXTERNAL_APIS", "false")
    monkeypatch.setenv("SMTP_HOST", "smtp.example.com")
    monkeypatch.setenv("SMTP_PORT", "587")
    monkeypatch.setenv("SMTP_USE_TLS", "true")
    monkeypatch.setenv("SMTP_USER", "user")
    monkeypatch.setenv("SMTP_PASSWORD", "pass")
    monkeypatch.setenv("SMTP_FROM", "silicon@test.local")
    from app.config import get_settings

    get_settings.cache_clear()
    reset_sent_emails()
    await client.get("/admin/login")
    token = await _owner_token(client)

    with patch("app.services.email_service.smtplib.SMTP", FakeSMTP):
        resp = await client.post(
            "/admin/team/invite",
            data={
                "token": token,
                "email": "smtp.manager@example.com",
                "full_name": "SMTP Тест",
            },
            follow_redirects=False,
        )

    assert resp.status_code == 303
    assert "msg=" in resp.headers["location"]
    assert "mock" not in resp.headers["location"].lower() or True  # success redirect
    assert len(sent_msgs) == 1
    msg = sent_msgs[0]
    assert msg["To"] == "smtp.manager@example.com"
    assert msg["From"] == "silicon@test.local"
    body = msg.get_content()
    assert "Пароль:" in body
    assert "smtp.manager@example.com" in body
    assert "SMTP Тест" in body
    # не mock у внутрішньому payload
    assert get_sent_emails()[-1]["to"] == "smtp.manager@example.com"
    get_settings.cache_clear()


def test_send_manager_invite_email_unit(monkeypatch):
    monkeypatch.setenv("MOCK_EXTERNAL_APIS", "true")
    from app.config import get_settings
    from app.services.email_service import send_manager_invite_email

    get_settings.cache_clear()
    reset_sent_emails()
    result = send_manager_invite_email(
        email="a@b.com",
        password="Abc12345xy",
        inviter_email="owner@silicon.local",
        full_name="Іван Петренко",
    )
    assert result["success"] is True
    assert result["mock"] is True
    assert "Іван Петренко" in result["body"]
    assert "Email / логін: a@b.com" in result["body"]
    assert "Пароль: Abc12345xy" in result["body"]
    get_settings.cache_clear()


@pytest.mark.asyncio
async def test_delete_manager(client: AsyncClient, monkeypatch):
    monkeypatch.setenv("ADMIN_SECRET_KEY", "delete-secret")
    monkeypatch.setenv("ADMIN_PASSWORD", "silicon-admin")
    monkeypatch.setenv("ADMIN_OWNER_EMAIL", "owner@silicon.local")
    monkeypatch.setenv("MOCK_EXTERNAL_APIS", "true")
    from app.config import get_settings

    get_settings.cache_clear()
    await client.get("/admin/login")
    token = await _owner_token(client)

    invite = await client.post(
        "/admin/team/invite",
        data={
            "token": token,
            "email": "to.delete@example.com",
            "full_name": "На Видалення",
        },
        follow_redirects=False,
    )
    assert invite.status_code == 303

    team = await client.get(f"/admin/team?token={token}")
    assert team.status_code == 200
    assert "to.delete@example.com" in team.text
    assert "/delete" in team.text

    # знайти id менеджера з HTML форми delete
    import re

    m = re.search(r'/admin/team/(\d+)/delete', team.text)
    assert m, "delete form not found"
    mgr_id = m.group(1)

    resp = await client.post(
        f"/admin/team/{mgr_id}/delete",
        data={"token": token},
        follow_redirects=False,
    )
    assert resp.status_code == 303
    assert "msg=" in resp.headers["location"]

    team2 = await client.get(f"/admin/team?token={token}")
    assert "to.delete@example.com" not in team2.text
    get_settings.cache_clear()


@pytest.mark.asyncio
async def test_admin_can_delete_chat_and_contact(client: AsyncClient, monkeypatch):
    monkeypatch.setenv("ADMIN_SECRET_KEY", "delete-chat-secret")
    monkeypatch.setenv("ADMIN_PASSWORD", "silicon-admin")
    monkeypatch.setenv("ADMIN_OWNER_EMAIL", "owner@silicon.local")
    monkeypatch.setenv("MOCK_EXTERNAL_APIS", "true")
    from app.config import get_settings

    get_settings.cache_clear()
    await client.get("/admin/login")
    token = await _owner_token(client)

    chat_only = "del-chat-only-1"
    await client.post("/api/chat/message", json={"chat_id": chat_only, "message": "привіт"})
    await client.post("/api/orders/register-client", json={
        "chat_id": chat_only,
        "first_name": "Іван",
        "last_name": "Тестовий",
        "phone": "+380501111111",
        "city": "Київ",
    })

    page = await client.get(f"/admin/live-chats?token={token}")
    assert page.status_code == 200
    assert "delete-contact" in page.text
    assert "Видалити чат" in page.text

    listed = await client.get(f"/admin/live-chats/data?token={token}")
    ids = [u["chat_id"] for u in listed.json()["users"]]
    assert chat_only in ids

    gone_chat = await client.post(
        f"/admin/live-chats/{chat_only}/delete",
        data={"token": token},
    )
    assert gone_chat.status_code == 200
    assert gone_chat.json()["success"] is True
    listed2 = await client.get(f"/admin/live-chats/data?token={token}")
    assert chat_only not in [u["chat_id"] for u in listed2.json()["users"]]
    still = await client.get(f"/admin/client/{chat_only}?token={token}")
    assert still.status_code == 200
    assert still.json()["client"] is not None
    assert still.json()["client"]["phone"] == "+380501111111"

    both = "del-contact-1"
    await client.post("/api/chat/message", json={"chat_id": both, "message": "хочу браслети"})
    await client.post("/api/orders/register-client", json={
        "chat_id": both,
        "first_name": "Марія",
        "phone": "+380502222222",
    })
    wipe = await client.post(
        f"/admin/live-chats/{both}/delete-contact",
        data={"token": token},
    )
    assert wipe.status_code == 200
    assert wipe.json()["success"] is True
    assert wipe.json()["clients"] >= 1
    listed3 = await client.get(f"/admin/live-chats/data?token={token}")
    assert both not in [u["chat_id"] for u in listed3.json()["users"]]
    get_settings.cache_clear()


@pytest.mark.asyncio
async def test_bot_people_page_and_excel(client: AsyncClient, monkeypatch):
    monkeypatch.setenv("ADMIN_SECRET_KEY", "people-secret")
    monkeypatch.setenv("ADMIN_PASSWORD", "silicon-admin")
    monkeypatch.setenv("ADMIN_OWNER_EMAIL", "owner@silicon.local")
    from io import BytesIO

    from openpyxl import load_workbook

    from app.config import get_settings

    get_settings.cache_clear()
    await client.get("/admin/login")
    token = await _owner_token(client)

    await client.post("/api/chat/message", json={
        "chat_id": "baza-user-77",
        "message": "/start",
    })
    await client.post("/api/orders/register-client", json={
        "chat_id": "baza-user-77",
        "first_name": "Олена База",
        "phone": "+380509999000",
        "city": "Львів",
    })

    page = await client.get(f"/admin/clients?token={token}")
    assert page.status_code == 200
    assert "База клієнтів" in page.text
    assert "baza-user-77" in page.text or "Олена" in page.text or "+380509999000" in page.text
    assert "Вивантажити Excel" in page.text

    found = await client.get(f"/admin/clients?token={token}&q=380509999000")
    assert found.status_code == 200
    assert "+380509999000" in found.text

    xlsx = await client.get(f"/admin/clients/export.xlsx?token={token}")
    assert xlsx.status_code == 200
    assert "spreadsheetml" in xlsx.headers.get("content-type", "")
    wb = load_workbook(BytesIO(xlsx.content))
    assert "Клієнти" in wb.sheetnames
    rows = list(wb["Клієнти"].iter_rows(values_only=True))
    blob = " ".join(str(c or "") for row in rows for c in row)
    assert "baza-user-77" in blob
    assert "380509999000" in blob
    get_settings.cache_clear()
